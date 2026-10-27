"""Тести: python -m pytest  (або python -m unittest discover tests)."""
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from merger.core import merge
from merger.nbu import get_rates
from merger.normalize import normalize_article, parse_price, parse_stock, stock_label
from merger.readers import Item, find_file, read_supplier

ROOT = Path(__file__).resolve().parent.parent


def item(sup, art, price, cur="UAH", stock=None, name="", brand=""):
    return Item(sup, normalize_article(art), art, name, brand, price, cur, stock, stock_label(stock), 1)


class NormalizeTests(unittest.TestCase):
    def test_articles(self):
        for raw in ["OC 90", "oc-90", "OC.90", " oc_90 "]:
            self.assertEqual(normalize_article(raw), "OC90")
        self.assertEqual(normalize_article(333114.0), "333114")
        self.assertEqual(normalize_article(None), "")

    def test_prices(self):
        cases = {"1 234,50": 1234.5, "$12.5": 12.5, "1.234,50": 1234.5, "1,234.50": 1234.5,
                 "109 грн": 109.0, 41: 41.0, "": None, "договірна": None, 0: None, "-5": None}
        for raw, exp in cases.items():
            self.assertEqual(parse_price(raw), exp, raw)

    def test_stock(self):
        self.assertEqual(parse_stock("є"), -1)
        self.assertEqual(parse_stock("+"), -1)
        self.assertEqual(parse_stock("-"), 0)
        self.assertEqual(parse_stock("під замовлення"), 0)
        self.assertEqual(parse_stock(">10"), 11)
        self.assertEqual(parse_stock(7), 7)
        self.assertIsNone(parse_stock(None))
        self.assertEqual(stock_label(11, ">10"), ">10")
        self.assertEqual(stock_label(0, ""), "немає")


class MergeTests(unittest.TestCase):
    def test_best_price_skips_out_of_stock(self):
        items = [item("A", "OC 90", 300, stock=5), item("B", "oc-90", 250, stock=0), item("C", "OC90", 280, stock=-1)]
        [p] = merge(items, {"UAH": 1})
        self.assertEqual(len(p.offers), 3)
        self.assertEqual(p.best.supplier, "C")      # B дешевше, але немає в наявності
        self.assertEqual(p.saving, 20.0)      # найдорожча 300 − найкраща 280

    def test_currency_conversion_and_markup(self):
        [p] = merge([item("A", "X1", 10, "USD")], {"UAH": 1, "USD": 41.5}, markup_percent=10)
        self.assertAlmostEqual(p.best.price_uah, 456.5)

    def test_nothing_in_stock_takes_cheapest(self):
        [p] = merge([item("A", "Z", 100, stock=0), item("B", "Z", 90, stock=0)], {"UAH": 1})
        self.assertEqual(p.best.supplier, "B")
        self.assertFalse(p.in_stock_anywhere)

    def test_duplicate_in_same_supplier_keeps_cheaper(self):
        [p] = merge([item("A", "Q", 100), item("A", "Q", 80)], {"UAH": 1})
        self.assertEqual(p.offers["A"].price_uah, 80)


class RatesTests(unittest.TestCase):
    def test_manual_has_priority_and_fetch_used_otherwise(self):
        rates, src = get_rates({"USD", "EUR", "UAH"}, {"USD": 41}, fetch=lambda c: 48.0)
        self.assertEqual(rates, {"UAH": 1.0, "USD": 41.0, "EUR": 48.0})
        self.assertEqual(src["EUR"], "НБУ")

    def test_fetch_error_is_clear(self):
        def boom(c):
            raise OSError("no internet")
        with self.assertRaises(RuntimeError) as cm:
            get_rates({"USD"}, {}, fetch=boom)
        self.assertIn("manual_rates", str(cm.exception))


class EndToEndTests(unittest.TestCase):
    def test_demo_prices(self):
        import merge_prices
        tmp = Path(tempfile.mkdtemp())
        try:
            shutil.copy(ROOT / "config.example.json", tmp / "config.json")
            shutil.copytree(ROOT / "input", tmp / "input")
            out = merge_prices.run(tmp, fetch=lambda c: {"USD": 41.6, "EUR": 48.3}[c])
            wb = load_workbook(out)
            rows = list(wb["Зведений прайс"].iter_rows(values_only=True))
            by_art = {normalize_article(r[0]): r for r in rows[1:]}
            self.assertEqual(len(by_art), 9)
            self.assertEqual(by_art["0986494104"][4], "МоторЛюкс")   # 31.9$ * 41.6 = 1327.04
            self.assertEqual(by_art["333114"][4], "Запчастина Плюс") # у МоторЛюкс немає в наявності
            self.assertEqual(by_art["BKR6E"][5], ">10")
            self.assertEqual(set(wb.sheetnames), {"Зведений прайс", "Усі позиції", "Звіт"})
        finally:
            shutil.rmtree(tmp)

    def test_missing_column_reported(self):
        cfg = {"name": "X", "header_row": 1, "currency": "UAH",
               "columns": {"article": "Нема такої", "price": "Ціна, грн"}}
        res = read_supplier(ROOT / "input" / "avtodetal_2026-10-27.xlsx", cfg)
        self.assertEqual(res.items, [])
        self.assertIn("Нема такої", res.errors[0])

    def test_find_file_by_pattern(self):
        self.assertEqual(find_file(ROOT / "input", "motorlux*.xlsx").name, "motorlux_oct.xlsx")
        self.assertIsNone(find_file(ROOT / "input", "nothing*.xlsx"))


if __name__ == "__main__":
    unittest.main()
