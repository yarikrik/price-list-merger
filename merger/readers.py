"""Читання прайсів постачальників з різною структурою (.xlsx / .csv)."""
from __future__ import annotations

import csv
import fnmatch
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .normalize import normalize_article, parse_price, parse_stock, stock_label

FIELDS = ("article", "name", "brand", "price", "stock")


@dataclass
class Item:
    supplier: str
    article: str        # нормалізований ключ для зведення
    article_raw: str    # як написано в прайсі
    name: str
    brand: str
    price: float        # у валюті постачальника
    currency: str
    stock: int | None
    stock_text: str     # як показувати наявність у звіті
    row: int


@dataclass
class ReadResult:
    supplier: str
    file: str
    items: list[Item]
    skipped: int        # рядки без артикула або ціни
    errors: list[str]


def _rows(path: Path, sheet: str | None):
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        text = None
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        dialect = csv.Sniffer().sniff(text.splitlines()[0] if text else ",", delimiters=";,\t")
        yield from csv.reader(text.splitlines(), dialect)
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    for row in ws.iter_rows(values_only=True):
        yield list(row)
    wb.close()


def _resolve_columns(spec: dict, header: list) -> dict[str, int]:
    """Колонка задається літерою ('C') або назвою в шапці ('Артикул'). Повертає індекси з 0."""
    names = [str(h).strip().lower() if h is not None else "" for h in header]
    out = {}
    for field, ref in spec.items():
        if field not in FIELDS or not ref:
            continue
        ref = str(ref).strip()
        if ref.lower() in names:  # спочатку шукаємо за назвою в шапці
            out[field] = names.index(ref.lower())
        elif ref.isalpha() and len(ref) <= 3 and ref.upper() == ref:
            out[field] = column_index_from_string(ref) - 1
        else:
            raise ValueError(f"колонку «{ref}» для поля {field} не знайдено в шапці: {header}")
    for need in ("article", "price"):
        if need not in out:
            raise ValueError(f"у налаштуваннях не вказана колонка для {need}")
    return out


def find_file(folder: Path, pattern: str) -> Path | None:
    matches = sorted(p for p in folder.iterdir() if p.is_file() and fnmatch.fnmatch(p.name.lower(), pattern.lower())
                     and not p.name.startswith("~$"))
    return matches[-1] if matches else None  # останній за назвою (наприклад, з найновішою датою)


def read_supplier(path: Path, cfg: dict) -> ReadResult:
    name = cfg["name"]
    header_row = int(cfg.get("header_row", 1))
    currency = cfg.get("currency", "UAH").upper()
    items, skipped, errors = [], 0, []
    rows = list(_rows(path, cfg.get("sheet")))
    if len(rows) < header_row:
        return ReadResult(name, path.name, [], 0, [f"у файлі менше {header_row} рядків"])
    try:
        cols = _resolve_columns(cfg["columns"], rows[header_row - 1])
    except ValueError as e:
        return ReadResult(name, path.name, [], 0, [str(e)])

    def get(row, field):
        i = cols.get(field)
        return row[i] if i is not None and i < len(row) else None

    for n, row in enumerate(rows[header_row:], start=header_row + 1):
        if not row or all(v in (None, "") for v in row):
            continue
        article = normalize_article(get(row, "article"))
        price = parse_price(get(row, "price"))
        if not article or price is None:
            skipped += 1
            continue
        items.append(Item(
            supplier=name, article=article, article_raw=str(get(row, "article")).strip(),
            name=str(get(row, "name") or "").strip(), brand=str(get(row, "brand") or "").strip(),
            price=price, currency=currency, stock=(qty := parse_stock(get(row, "stock"))),
            stock_text=stock_label(qty, get(row, "stock")), row=n))
    return ReadResult(name, path.name, items, skipped, errors)
