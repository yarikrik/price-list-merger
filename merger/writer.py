"""Формування підсумкового Excel-файлу."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .core import Position
from .readers import ReadResult

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
BEST_FILL = PatternFill("solid", fgColor="C6EFCE")
NO_STOCK_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0.00 "грн"'


def _header(ws, titles, widths):
    ws.append(titles)
    for i, (cell, w) in enumerate(zip(ws[1], widths), start=1):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"


def write_report(path: Path, positions: list[Position], suppliers: list[str], results: list[ReadResult],
                 rates: dict, rate_sources: dict, markup: float = 0) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Зведений прайс"
    titles = ["Артикул", "Назва", "Бренд", "Найкраща ціна", "Постачальник", "Наявність", "Економія"] + \
             [f"{s}" for s in suppliers]
    _header(ws, titles, [16, 42, 14, 15, 18, 11, 12] + [14] * len(suppliers))

    for p in positions:
        best = p.best
        row = [p.article_raw, p.name, p.brand, best.price_uah, best.supplier, best.stock_text,
               p.saving or None]
        for s in suppliers:
            o = p.offers.get(s)
            row.append(o.price_uah if o else None)
        ws.append(row)
        r = ws.max_row
        for c in ws[r]:
            c.border = BORDER
        ws.cell(r, 4).number_format = MONEY
        ws.cell(r, 7).number_format = MONEY
        ws.cell(r, 4).font = Font(bold=True)
        if not p.in_stock_anywhere:
            for c in ws[r][:7]:
                c.fill = NO_STOCK_FILL
        for j, s in enumerate(suppliers):
            cell = ws.cell(r, 8 + j)
            cell.number_format = MONEY
            o = p.offers.get(s)
            if o and o.supplier == best.supplier:
                cell.fill = BEST_FILL
                cell.font = Font(bold=True)
            elif o and not o.available:
                cell.font = Font(color="999999", italic=True)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(titles))}{max(ws.max_row, 1)}"

    # Усі позиції як є — для перевірки
    raw = wb.create_sheet("Усі позиції")
    _header(raw, ["Постачальник", "Артикул (як у прайсі)", "Ключ зведення", "Назва", "Бренд",
                  "Ціна (оригінал)", "Валюта", "Ціна, грн", "Наявність", "Рядок у файлі"],
            [18, 18, 16, 40, 14, 14, 9, 14, 11, 10])
    for res in results:
        for it in res.items:
            raw.append([it.supplier, it.article_raw, it.article, it.name, it.brand, it.price, it.currency,
                        round(it.price * rates[it.currency] * (1 + markup / 100), 2), it.stock_text, it.row])
            raw.cell(raw.max_row, 8).number_format = MONEY
    raw.auto_filter.ref = f"A1:J{max(raw.max_row, 1)}"

    # Звіт
    rep = wb.create_sheet("Звіт")
    rep.column_dimensions["A"].width = 34
    rep.column_dimensions["B"].width = 26
    rep.column_dimensions["C"].width = 14
    rep.column_dimensions["D"].width = 14
    lines = [("Звіт зведення прайсів", None), ("Сформовано", datetime.now().strftime("%d.%m.%Y %H:%M")),
             ("Унікальних позицій", len(positions)),
             ("Позицій у кількох постачальників", sum(1 for p in positions if len(p.offers) > 1)),
             ("Націнка", f"{markup}%"), ("", None), ("Курси валют", None)]
    for cur, rate in sorted(rates.items()):
        if cur != "UAH":
            lines.append((f"{cur} → UAH", f"{rate:.4f} ({rate_sources.get(cur, '')})"))
    lines += [("", None), ("Постачальник", "Файл")]
    for line in lines:
        rep.append(list(line))
    start = rep.max_row
    rep.cell(start, 3, "Позицій")
    rep.cell(start, 4, "Пропущено")
    for res in results:
        rep.append([res.supplier, res.file, len(res.items), res.skipped])
        for e in res.errors:
            rep.append(["", f"⚠ {e}"])
    rep["A1"].font = Font(bold=True, size=14)
    for r in (7, start):
        for c in rep[r]:
            c.font = Font(bold=True)
    wb.save(path)
