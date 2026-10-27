"""Генерує 4 демо-прайси з різною структурою в папку input/ (для перевірки та скриншотів)."""
import csv
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
INP = ROOT / "input"
INP.mkdir(exist_ok=True)

# 1. АвтоДеталь — звичайна таблиця, ціни в грн
wb = Workbook(); ws = wb.active
ws.append(["Найменування", "Бренд", "Артикул", "Ціна, грн", "Наявність"])
for r in [
    ["Фільтр масляний", "KNECHT", "OC 90", 289, "є"],
    ["Фільтр повітряний", "MANN", "C 2774/3", 455, 12],
    ["Колодки гальмівні передні", "BOSCH", "0 986 494 104", 1390, 4],
    ["Свічка запалювання", "NGK", "BKR6E", 118, "є"],
    ["Ремінь ГРМ", "GATES", "5578XS", 1120, 0],
    ["Амортизатор передній", "KYB", "333114", 2350, 2],
    ["Фільтр салону", "MANN", "CU 2939", 410, "немає"],
]:
    ws.append(r)
wb.save(INP / "avtodetal_2026-10-27.xlsx")

# 2. МоторЛюкс — шапка на 5 рядків, ціни в USD, колонки за літерами
wb = Workbook(); ws = wb.active
ws.append(["ТОВ «МоторЛюкс» — оптовий прайс"]); ws.append(["Тел.: +38 (000) 000-00-00"])
ws.append(["Ціни вказані в USD"]); ws.append(["Дата: 27.10.2026"]); ws.append([])
ws.append(["Код", "Опис", "Виробник", "Ціна USD", "Залишок"])
for r in [
    ["OC-90", "Масляний фільтр", "Knecht", 6.4, 25],
    ["0986494104", "Brake pads front", "Bosch", 31.9, 10],
    ["BKR6E", "Spark plug", "NGK", 3.1, 100],
    ["5578XS", "Timing belt", "Gates", 24.5, 3],
    ["333114", "Shock absorber", "KYB", 59.0, 0],
    ["K015578XS", "Комплект ГРМ", "Gates", 89.0, 1],
]:
    ws.append(r)
wb.save(INP / "motorlux_oct.xlsx")

# 3. ЄвроПартс — CSV з «;», ціни в EUR з комою, наявність +/-
with open(INP / "europarts_price.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f, delimiter=";")
    w.writerow(["SKU", "Brand", "Name", "Price", "Stock"])
    for r in [
        ["c2774/3", "MANN", "Air filter", "8,90", "+"],
        ["CU2939", "MANN", "Cabin filter", "7,20", "+"],
        ["0 986 494 104", "BOSCH", "Brake pad set", "29,50", "-"],
        ["oc90", "KNECHT", "Oil filter", "5,95", "+"],
    ]:
        w.writerow(r)

# 4. Запчастина Плюс — шапка на 2 рядки, ціни текстом «1 250 грн», склад «>10»
wb = Workbook(); ws = wb.active
ws.append(["Запчастина Плюс"]); ws.append(["Оновлено щоденно"])
ws.append(["Товар", "Марка", "Кат. номер", "Роздріб", "Склад"])
for r in [
    ["Свічка запалювання", "NGK", "BKR-6E", "109 грн", ">10"],
    ["Колодки гальмівні", "Bosch", "0986494104", "1 345 грн", "3"],
    ["Ремінь ГРМ", "Gates", "5578 XS", "1 089 грн", "під замовлення"],
    ["Амортизатор", "KYB", "333114", "2 290 грн", "1"],
    ["Акумулятор 60Ah", "Varta", "560408054", "3 950 грн", "5"],
    ["", "", "", "", ""],
]:
    ws.append(r)
wb.save(INP / "zplus_27.10.xlsx")
print("Демо-прайси створено в", INP)
